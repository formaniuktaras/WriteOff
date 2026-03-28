from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import DocumentType, Nomenclature, Service, Unit


class ExcelDictionaryImportService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def preview(self, file_path: Path, limit: int = 10) -> list[list[str]]:
        wb = load_workbook(file_path, read_only=True)
        ws = wb.active
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=limit, values_only=True):
            rows.append([str(c) if c is not None else "" for c in row])
        return rows

    def import_rows(self, dictionary: str, rows: list[dict]) -> int:
        mapper = {
            "units": Unit,
            "services": Service,
            "nomenclature": Nomenclature,
            "document_types": DocumentType,
        }
        model = mapper[dictionary]
        count = 0
        for row in rows:
            entity = self.db.query(model).filter_by(code=row["code"]).first()
            if not entity:
                entity = model(code=row["code"], name=row["name"])
                self.db.add(entity)
            else:
                entity.name = row["name"]
            if dictionary == "units":
                setattr(entity, "parent_id", row.get("parent_id"))
            count += 1
        self.db.commit()
        return count
